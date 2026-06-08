import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('MAU_02_Template.xlsx')
print('Sheets:', wb.sheetnames)
ws = wb[wb.sheetnames[0]]
print('Dimensions:', ws.dimensions)
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True)):
    print(f'Row {i+1}: {row}')